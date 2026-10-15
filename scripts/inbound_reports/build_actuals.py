#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""(GS) 인바운드 실적_업체별,사업장별 — 지정 연도 파일 생성.
2024 양식(헤더/열구조/서식)을 유지하되, 데이터 행(업체×국가)은 해당 연도 전체 업체로
새로 생성한다(누락 0). 업체는 RNs 내림차순, 업체 내 국가도 내림차순(기타 말미).

사용:
  python3 scripts/inbound_reports/build_actuals.py --year 2025
  python3 scripts/inbound_reports/build_actuals.py --year 2026 --max-month 09
  python3 scripts/inbound_reports/build_actuals.py --year 2025 2026
"""
import os, re, sys, argparse, shutil, collections
from copy import copy
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import scan_inbound

HOME = os.path.expanduser('~')
DEF_TMPL = os.path.join(HOME, 'Downloads', '(GS)2024년 인바운드 실적_업체별,사업장별.xlsx')

def nsp(s): return re.sub(r'\s', '', str(s or ''))

_COL_PAIRS = [('소노캄비발디','캄비발디'),('소노펫비발디','소노펫비발디'),('소노펠리체빌리지비발디','소펠빌리지'),
    ('소노펠리체비발디','소펠비발디'),('비발디','비발디'),
    ('소노펠리체빌리지델피노','소펠빌델피노'),('소노펠리체델피노','소펠델피노'),
    ('소노문델피노','델피노(문)'),('소노벨델피노','델피노(벨)'),('소노캄델피노','델피노(캄)'),
    ('양평','양평'),('양양','양양'),('삼척','삼척'),('경주','경주'),('청송','청송'),('단양','단양'),
    ('천안','천안'),('변산','변산'),('거제','거제'),('여수','여수'),('진도','진도'),
    ('소노캄제주','제주(캄)'),('소노벨제주','제주(벨)'),('소노문제주','제주(벨)'),
    ('고양','고양'),('해운대','소노문해운대')]
_SHEET_PAIRS = [('소노캄비발디','비발디(펫)'),('소노펫비발디','비발디(펫)'),('소노펠리체빌리지비발디','비발디(펠리체)'),
    ('소노펠리체비발디','비발디(펠리체)'),('비발디','비발디(문벨)'),
    ('소노문델피노','문델피노'),('소노벨델피노','벨델피노'),('소노캄델피노','캄델피노'),('소노펠리체델피노','벨델피노'),
    ('양평','양평'),('양양','양양'),('삼척','삼척'),('경주','경주'),('청송','청송'),('단양','단양'),
    ('천안','천안'),('변산','변산'),('거제','거제'),('여수','여수'),('진도','진도'),
    ('소노벨제주','벨제주'),('고양','고양'),('해운대','해운대')]
def _match(pn, pairs):
    s = nsp(pn)
    for k, v in pairs:
        if k in s: return nsp(v)
    return None

def _nat_order(d):
    return sorted(d, key=lambda k: (k == '기타', -d.get(k, 0)))

def build(year, max_month, template, outdir):
    lines = scan_inbound(year, max_month)
    # scope='Total'(사업장별 열) / scope=sheetkey(월별 열)
    # cell[scope][(업체,국가)][colkey or 'TOTAL'] = [rns, rev]
    cell = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(lambda: [0, 0.0])))
    for ln in lines:
        pcol = _match(ln['prop'], _COL_PAIRS); psheet = _match(ln['prop'], _SHEET_PAIRS)
        for bk in (ln['agent'], 'Total'):
            for ck in (ln['nat'], 'Total'):
                for col in ([pcol, 'TOTAL'] if pcol else ['TOTAL']):
                    e = cell['Total'][(bk, ck)][nsp(col)]; e[0] += ln['rns']; e[1] += ln['rev']
                if psheet:
                    for col in (ln['month'], 'TOTAL'):
                        e = cell[psheet][(bk, ck)][nsp(col)]; e[0] += ln['rns']; e[1] += ln['rev']

    out = os.path.join(outdir, f'(GS){year}년 인바운드 실적_업체별,사업장별.xlsx')
    shutil.copy(template, out)
    wb = openpyxl.load_workbook(out)

    def gen_rows(scope):
        """(업체, 국가, kind) 순서 리스트. kind: grand/agent_total/agent_nat"""
        d = cell.get(scope, {})
        grand_nats = {ck: d[('Total', ck)]['TOTAL'][0] for (bk, ck) in d if bk == 'Total' and ck != 'Total'}
        rows = [('Total', 'Total', 'grand')] + [('Total', ck, 'grand') for ck in _nat_order(grand_nats) if grand_nats[ck]]
        agents = {bk: d[(bk, 'Total')]['TOTAL'][0] for (bk, ck) in d if ck == 'Total' and bk != 'Total'}
        for ag in sorted(agents, key=lambda a: -agents[a]):
            if not agents[ag]: continue
            nats = {ck: d[(ag, ck)]['TOTAL'][0] for (bk, ck) in d if bk == ag and ck != 'Total'}
            rows.append((ag, 'Total', 'agent_total'))
            rows += [(ag, ck, 'agent_nat') for ck in _nat_order(nats) if nats[ck]]
        return rows

    def process(ws, scope):
        # 열 매핑 (row5): RNs열 → colkey
        cols = [(ci + 1, nsp(v)) for ci, v in enumerate([c.value for c in ws[5]])
                if v not in (None, '') and str(v).strip() not in ('구 분', '구분')]
        is_total = (scope == 'Total')
        # 스타일 템플릿 캡처 (삭제 전)
        styref = {}  # kind → {col: _style}
        for r in range(7, ws.max_row + 1):
            b = ws.cell(r, 2).value; c = ws.cell(r, 3).value
            if b in (None, '') or c in (None, ''): continue
            b = str(b).strip(); c = str(c).strip()
            kind = 'grand' if b == 'Total' else ('agent_total' if c == 'Total' else 'agent_nat')
            if kind not in styref:
                styref[kind] = {cc: copy(ws.cell(r, cc)._style) for cc in range(2, ws.max_column + 1)}
            if len(styref) == 3: break
        if 'agent_total' not in styref: styref['agent_total'] = styref.get('grand', {})
        if 'agent_nat' not in styref: styref['agent_nat'] = styref.get('agent_total', styref.get('grand', {}))
        rows = gen_rows(scope)
        d = cell.get(scope, {})
        last = ws.max_row
        # 데이터행 삭제 후 재작성
        if last >= 7:
            ws.delete_rows(7, last - 6)
        for i, (bk, ck, kind) in enumerate(rows):
            r = 7 + i
            st = styref.get(kind, {})
            for cc in range(2, ws.max_column + 1):
                if cc in st: ws.cell(r, cc)._style = copy(st[cc])
            ws.cell(r, 2).value = bk
            ws.cell(r, 3).value = ck
            rowd = d.get((bk, ck), {})
            for rc, ckey in cols:
                key = ckey if (is_total or ckey == 'TOTAL') else ckey.zfill(2)
                rns, rev = rowd.get(key, [0, 0.0])
                adr = (rev * 1e6 / rns / 1000) if rns else 0
                ws.cell(r, rc).value = rns if rns else 0
                ws.cell(r, rc + 1).value = round(adr, 1) if rns else 0
                ws.cell(r, rc + 2).value = round(rev, 1) if rev else 0
        # 템플릿의 잔존 숨김행 제거 — 데이터행 전부 보이게 (대만 등 숨김 방지)
        for rr in range(7, max(ws.max_row, 7 + len(rows)) + 1):
            if rr in ws.row_dimensions:
                ws.row_dimensions[rr].hidden = False

    process(wb['Total'], 'Total')
    for sh in wb.sheetnames:
        if sh == 'Total': continue
        process(wb[sh], nsp(sh))
    # 제목 연도
    for sh in wb.sheetnames:
        ws = wb[sh]
        for r in range(1, 4):
            for c in range(1, 6):
                v = ws.cell(r, c).value
                if isinstance(v, str) and '24년' in v:
                    ws.cell(r, c).value = v.replace('24년', f'{str(year)[2:]}년')
    wb.save(out)
    g = cell['Total'][('Total', 'Total')]['TOTAL'][0]
    print(f'{year}: 라인 {len(lines):,} / Total RNs={g:,} → {out}')

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--year', nargs='+', type=int, required=True)
    ap.add_argument('--max-month', default='12')
    ap.add_argument('--template', default=DEF_TMPL)
    ap.add_argument('--outdir', default=os.path.join(HOME, 'Downloads'))
    a = ap.parse_args()
    for y in a.year:
        build(y, a.max_month, a.template, a.outdir)
