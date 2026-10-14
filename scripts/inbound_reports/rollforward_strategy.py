#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""(인바운드) 전략방안 — 연도 1칸씩 롤포워드 + 실적 재집계, 피벗 제거 경량화(양식 보존).

BASE_YEAR 기준으로 총괄/사업장별 시트의 연도열을 한 칸씩 당겨 채운다.
  '2025 실적' 슬롯(RN col17) → BASE_YEAR
  '2024' 슬롯(23) → BASE_YEAR-1
  '2023' 슬롯(29) → BASE_YEAR-2
  '2019' 슬롯(35) → BASE_YEAR-3
각 연도는 raw_db 집계코드58 인바운드 재집계(최신연도는 --max-month 까지). 피벗/캐시 제거.

사용:
  python3 scripts/inbound_reports/rollforward_strategy.py --base-year 2026 --max-month 09
내년:
  python3 scripts/inbound_reports/rollforward_strategy.py --base-year 2027 --max-month 09
옵션: --template(기본 Downloads 원본), --out(기본 Downloads '_실적반영.xlsx')
"""
import os, re, sys, argparse, collections, zipfile, shutil
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import scan_inbound

HOME = os.path.expanduser('~')
DEF_TMPL = os.path.join(HOME, 'Downloads', '(인바운드)2026 전략방안_112016.xlsx')
PREFIX2SHEET = {'01':'비발디','02':'비발디','03':'비발디','04':'비발디','05':'비발디','06':'비발디',
    '07':'양평','08':'델피노','09':'양양','10':'삼척','11':'단양','12':'경주','13':'청송',
    '14':'천안','15':'변산','16':'여수','17':'거제','18':'진도','21':'고양','22':'해운대'}
SHEETS = ['총괄(사업장별 ▶)','비발디','양평','델피노','양양','삼척','단양','경주','청송','천안','변산','여수','거제','진도','고양','해운대']

def strip_pivots(path):
    tmp = path + '.tmp'; zin = zipfile.ZipFile(path)
    isp = lambda n: ('pivotTable' in n) or ('pivotCache' in n)
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for n in zin.namelist():
        if isp(n): continue
        data = zin.read(n)
        if n == '[Content_Types].xml':
            data = re.sub(r'<Override[^>]*pivot(Table|Cache)[^>]*/>', '', data.decode('utf8','replace')).encode('utf8')
        elif n == 'xl/workbook.xml':
            data = re.sub(r'<pivotCaches>.*?</pivotCaches>', '', data.decode('utf8','replace'), flags=re.S).encode('utf8')
        elif n.endswith('.rels'):
            data = re.sub(r'<Relationship[^>]*pivot(Table|Cache)[^>]*/>', '', data.decode('utf8','replace')).encode('utf8')
        zout.writestr(n, data)
    zout.close(); zin.close(); shutil.move(tmp, path)

def build(base_year, max_month, template, out):
    year_col = {base_year:17, base_year-1:23, base_year-2:29, base_year-3:35}  # RN 시작열(1-based)
    agg = {}  # year → cell[(scope,월,dim,name)] = [rns,rev]
    for y in year_col:
        A = collections.defaultdict(lambda: [0, 0.0])
        for ln in scan_inbound(y, max_month if y == base_year else '12'):
            pref = ln['prop'][:2] if ln['prop'][:2].isdigit() else ''
            sheet = PREFIX2SHEET.get(pref)
            for scope in (['총괄'] + ([sheet] if sheet else [])):
                for mmk in (ln['month'], '년간'):
                    for dim, name in (('계','계'), ('국가', ln['nat']), ('Agent', ln['agent'])):
                        e = A[(scope, mmk, dim, name)]; e[0] += ln['rns']; e[1] += ln['rev']
        agg[y] = A
        print(f'  집계 {y}: 년간 계 RNs={A[("총괄","년간","계","계")][0]:,}')

    shutil.copy(template, out)
    wb = openpyxl.load_workbook(out)
    def g(y, scope, mm, dim, name): return agg[y].get((scope, mm, dim, name), [0, 0.0])
    for sh in SHEETS:
        if sh not in wb.sheetnames: continue
        ws = wb[sh]; scope = '총괄' if sh.startswith('총괄') else sh
        for y, base in year_col.items():
            for col in range(base, base+6): ws.cell(3, col).value = y
            ws.cell(4, base).value = f'{y}'
        cur_mm = cur_dim = None; blk = {}
        for r in range(7, ws.max_row+1):
            b = ws.cell(r,2).value; c = ws.cell(r,3).value; d = ws.cell(r,4).value
            if b not in (None,''):
                s = str(b).strip(); cur_mm = '년간' if s=='년간' else (s.zfill(2) if s.isdigit() else cur_mm)
            if c not in (None,''):
                cc = str(c).strip(); cur_dim = cc if cc in ('계','국가','Agent') else cur_dim
            if cur_mm is None: continue
            is_total = c is not None and str(c).strip()=='계'
            if is_total: name, dim = '계', '계'
            elif cur_dim in ('국가','Agent') and d not in (None,''): name, dim = str(d).strip(), cur_dim
            else: continue
            for y, base in year_col.items():
                rns, rev = g(y, scope, cur_mm, dim, name)
                adr = (rev*1e6/rns/1000) if rns else 0
                ws.cell(r, base).value   = rns if rns else 0
                ws.cell(r, base+2).value = round(adr,1) if rns else 0
                ws.cell(r, base+4).value = round(rev,1) if rev else 0
                if is_total: blk[y] = (rns, rev)
                else:
                    trns, trev = blk.get(y, (0,0))
                    tadr = (trev*1e6/trns/1000) if trns else 0
                    ws.cell(r, base+1).value = round(rns/trns,3) if trns else 0
                    ws.cell(r, base+3).value = round(adr/tadr,2) if tadr else 0
                    ws.cell(r, base+5).value = round(rev/trev,3) if trev else 0
    wb.save(out)
    strip_pivots(out)
    print('저장(피벗 제거):', out)

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--base-year', type=int, required=True)
    ap.add_argument('--max-month', default='09')
    ap.add_argument('--template', default=DEF_TMPL)
    ap.add_argument('--out', default=os.path.join(HOME, 'Downloads', '(인바운드) 전략방안_실적반영.xlsx'))
    a = ap.parse_args()
    build(a.base_year, a.max_month, a.template, a.out)
