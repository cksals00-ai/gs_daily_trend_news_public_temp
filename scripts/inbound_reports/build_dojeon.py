#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""★★도전목표_2026_전략방안 — 매일 실적(인바운드 온북, raw_db 코드58) 반영.

산출물 3종:
  1) docs/inbound-dojeon.html  의 inline `const D=...,PROPS=[...]`  (달성률 대시보드)
  2) docs/dojeon_strategy_actual.xlsx                              (대시보드 다운로드 버튼용, 호스팅)
  3) ~/Downloads/★★도전목표_2026_전략방안_실적반영.xlsx            (사용자 보고용 사본)

원본 양식(총괄+18사업장)에 「2026 Actual」열을 채우며, 템플릿은 피벗을 값으로 치환한 경량본.
build.py/run.sh 에서 호출하면 매일 자동 갱신된다.
"""
import os, re, sys, json, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _common as C
import openpyxl, warnings; warnings.filterwarnings('ignore')

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
TEMPLATE = os.path.join(HERE, 'templates', 'dojeon_2026.xlsx')
TARGETS  = json.load(open(os.path.join(HERE, 'dojeon_targets.json'), encoding='utf-8'))
DOCS_HTML = os.path.join(REPO, 'docs', 'inbound-dojeon.html')
DOCS_XLSX = os.path.join(REPO, 'docs', 'dojeon_strategy_actual.xlsx')
DL_XLSX   = os.path.expanduser('~/Downloads/★★도전목표_2026_전략방안_실적반영.xlsx')
YEAR = 2026

# ── 사업장 시트/열 레이아웃 (고양 구조와 동일) ──────────────────────────
# actual 롤: 2026→23, 2025→29, 2024→35, 2023→41 (각 그룹 6열: RN/비중 ADR/지수 Rev/비중)
YEAR_COL = {2026: 23, 2025: 29, 2024: 35, 2023: 41}
PREFIX2SHEET = {'01':'비발디','02':'비발디','03':'비발디','04':'비발디','05':'비발디','06':'비발디',
 '07':'양평','08':'델피노','09':'양양','10':'삼척','11':'단양','12':'경주','13':'청송','14':'천안',
 '15':'변산','16':'여수','17':'거제','18':'진도','19':'벨제주','20':'캄제주','21':'고양','22':'해운대','23':'남해'}
SHEETS = ['총괄(사업장별 ▶)','고양','해운대','비발디','단양','여수','델피노','삼척','벨제주','양평',
          '캄제주','천안','변산','거제','청송','양양','경주','진도','남해']

def _short(cp):
    s = re.sub(r'^\s*\d+\.\s*', '', cp or '')
    if '비발디' in s: return '비발디'
    if '소노캄 제주' in s or '캄제주' in s: return '캄제주'
    if '제주' in s: return '벨제주'
    for key in ['고양','해운대','단양','여수','델피노','삼척','천안','변산','거제',
                '청송','양양','경주','진도','남해','양평']:
        if key in s: return key
    return None


def main():
    # ── 실적 집계 ─────────────────────────────────────────────
    lines = list(C.scan_inbound(YEAR, '12'))
    # (1) 대시보드용: 사업장 short × 월
    act = collections.defaultdict(lambda: collections.defaultdict(lambda: [0, 0.0]))
    # (2) 엑셀용: year × (scope, 월, dim, name)  — scope=총괄 또는 사업장시트명
    xl = {y: collections.defaultdict(lambda: [0, 0.0]) for y in YEAR_COL}
    for ln in lines:
        sh = _short(ln['cprop'] or ln['prop'])
        if sh:
            for mk in (ln['month'], '계'):
                e = act[sh][mk]; e[0] += ln['rns']; e[1] += ln['rev']
    for y in YEAR_COL:
        for ln in C.scan_inbound(y, '12'):
            pref = (ln['cprop'] or ln['prop'])[:2]
            sheet = PREFIX2SHEET.get(pref if pref.isdigit() else '')
            for scope in (['총괄'] + ([sheet] if sheet else [])):
                for mmk in (ln['month'], '년간'):
                    for dim, name in (('계', '계'), ('국가', ln['nat']), ('Agent', ln['agent'])):
                        e = xl[y][(scope, mmk, dim, name)]; e[0] += ln['rns']; e[1] += ln['rev']

    # ── 대시보드 데이터 ────────────────────────────────────────
    D = {}
    for sh in set(TARGETS) | set(act):
        D[sh] = {'target': TARGETS.get(sh, {}),
                 'actual': {mk: act[sh][mk][0] for mk in act[sh]},
                 'actual_rev': {mk: round(act[sh][mk][1]) for mk in act[sh]}}
    tt = {'계': sum(TARGETS[s].get('계', 0) for s in TARGETS)}
    for m in range(1, 13):
        tt[str(m)] = sum(TARGETS[s].get(str(m), 0) for s in TARGETS)
    ta = collections.defaultdict(int)
    for sh in act:
        for mk, v in act[sh].items(): ta[mk] += v[0]
    D['전체'] = {'target': tt, 'actual': dict(ta), 'actual_rev': {}}
    PROPS = [s for s in ['고양','비발디','해운대','단양','여수','델피노','삼척','벨제주','양평',
             '청송','캄제주','경주','거제','양양','천안','변산','진도','남해'] if s in D]
    block = (f"const D={json.dumps(D, ensure_ascii=False)},"
             f"PROPS={json.dumps(PROPS, ensure_ascii=False)};")
    html = open(DOCS_HTML, encoding='utf-8').read()
    html2, n = re.subn(r'const D=.*?PROPS=\[[^\]]*\];', block, html, count=1, flags=re.S)
    if n != 1:
        raise SystemExit('ERR: dashboard inline data 패턴 미발견')
    open(DOCS_HTML, 'w', encoding='utf-8').write(html2)
    print(f'대시보드 갱신: 전체 도전 {tt["계"]:,} / 실적 {ta["계"]:,} '
          f'({ta["계"]/tt["계"]*100:.1f}%)')

    # ── 엑셀 채움 ─────────────────────────────────────────────
    def g(y, scope, mm, dim, name): return xl[y].get((scope, mm, dim, name), [0, 0.0])
    wb = openpyxl.load_workbook(TEMPLATE)
    for sh in SHEETS:
        if sh not in wb.sheetnames:
            print('  시트없음', sh); continue
        ws = wb[sh]; scope = '총괄' if sh.startswith('총괄') else sh
        for year, base in YEAR_COL.items():
            for col in range(base, base + 6):
                if ws.cell(3, col).value not in (None, ''):
                    ws.cell(3, col).value = year
            ws.cell(4, base).value = (f'{year} Actual' if year == 2026 else f'{year}')
        cur_mm = cur_dim = None; blk = {}
        for r in range(7, ws.max_row + 1):
            b = ws.cell(r, 2).value; c = ws.cell(r, 3).value; d = ws.cell(r, 4).value
            if b not in (None, ''):
                s = str(b).strip()
                cur_mm = '년간' if s == '년간' else (s.zfill(2) if s.isdigit() else cur_mm)
            if c not in (None, ''):
                cc = str(c).strip(); cur_dim = cc if cc in ('계', '국가', 'Agent') else cur_dim
            if cur_mm is None: continue
            is_total = c is not None and str(c).strip() == '계'
            if is_total: name, dim = '계', '계'
            elif cur_dim in ('국가', 'Agent') and d not in (None, ''):
                name, dim = str(d).strip(), cur_dim
            else: continue
            for year, base in YEAR_COL.items():
                rns, rev = g(year, scope, cur_mm, dim, name)
                adr = (rev * 1e6 / rns / 1000) if rns else 0
                ws.cell(r, base).value = rns if rns else 0
                ws.cell(r, base + 2).value = round(adr, 1) if rns else 0
                ws.cell(r, base + 4).value = round(rev, 1) if rev else 0
                if is_total: blk[year] = (rns, rev)
                else:
                    trns, trev = blk.get(year, (0, 0))
                    tadr = (trev * 1e6 / trns / 1000) if trns else 0
                    ws.cell(r, base + 1).value = round(rns / trns, 3) if trns else 0
                    ws.cell(r, base + 3).value = round(adr / tadr, 2) if tadr else 0
                    ws.cell(r, base + 5).value = round(rev / trev, 3) if trev else 0
    wb.save(DOCS_XLSX)
    try:
        import shutil; shutil.copyfile(DOCS_XLSX, DL_XLSX)
    except Exception as e:
        print('  Downloads 사본 실패:', e)
    print(f'엑셀 저장: {DOCS_XLSX}')
    print(f'         : {DL_XLSX}')


if __name__ == '__main__':
    main()
