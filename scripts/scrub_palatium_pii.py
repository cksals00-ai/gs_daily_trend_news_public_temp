#!/usr/bin/env python3
"""
팔라티움 예약정보조회 xlsx 개인정보 정리
==========================================
data/palatium_db/예약정보조회*.xlsx 를 덮어써서 개인정보를 제거한다.
- '비고', '핸드폰번호', '연락처' 열 삭제 (매번 수작업으로 지우던 3개 열)
- 고객·직원 이름 열(투숙객명·예약자명·영업직원·확인자·생성자·수정자)은 값만 비움.
  열을 지우면 위치가 밀려 위치 기반 파서(parse_palatium_db, build_palatium_segment_2026)가 깨지므로 남겨둔다.
- 이미 정리된 파일은 건드리지 않음(재실행 안전).
- parse_palatium_db.py 가 파싱 전에 자동 호출하므로, 커밋되는 원본에 남지 않는다.

단독 실행:  python3 scripts/scrub_palatium_pii.py
"""
import fs_utils  # macOS NFD→NFC 유니코드 정규화
import logging
import sys
from pathlib import Path

import openpyxl

PROJECT_DIR = Path(__file__).resolve().parent.parent
PALATIUM_DIR = PROJECT_DIR / "data" / "palatium_db"
DROP_COLUMNS = {"비고", "핸드폰번호", "연락처"}
BLANK_COLUMNS = {"투숙객명", "예약자명", "영업직원", "확인자", "생성자", "수정자"}

log = logging.getLogger("scrub_palatium_pii")


def _header_row(ws):
    """'상태'·'예약번호' 가 있는 헤더 행 번호(1-base). 없으면 None."""
    for r in range(1, 11):
        vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if "상태" in vals and "예약번호" in vals:
            return r, vals
    return None, None


def scrub_file(path: Path) -> int:
    """수정한 열 개수 반환(삭제 + 값 비움). 이미 정리된 파일은 0."""
    wb = openpyxl.load_workbook(str(path))
    ws = wb.active
    hr, header = _header_row(ws)
    if hr is None:
        log.warning(f"  헤더 못 찾음 — 스킵: {path.name}")
        return 0
    blanked = 0
    for i, h in enumerate(header):
        if h not in BLANK_COLUMNS:
            continue
        hit = False
        for (cell,) in ws.iter_rows(min_row=hr + 1, min_col=i + 1, max_col=i + 1):
            if cell.value not in (None, ""):
                cell.value = None
                hit = True
        blanked += hit
    targets = [i + 1 for i, h in enumerate(header) if h in DROP_COLUMNS]
    for col in sorted(targets, reverse=True):  # 뒤에서부터 지워야 인덱스가 안 밀림
        ws.delete_cols(col)
    if blanked or targets:
        wb.save(str(path))
    return blanked + len(targets)


def scrub_all(folder: Path = PALATIUM_DIR) -> int:
    n_files = 0
    for path in sorted(folder.glob("예약정보조회*.xlsx")):
        if path.name.startswith("~$"):
            continue
        n = scrub_file(path)
        if n:
            n_files += 1
            log.info(f"  개인정보 열 {n}개 정리: {path.name}")
    return n_files


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else PALATIUM_DIR
    cnt = scrub_all(folder)
    log.info(f"정리 완료: {cnt}개 파일 수정")
