import importlib.util
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

MODULE = Path(__file__).resolve().parents[1] / "import_onbook.py"


def row(id=1, **extra):
    value = {"Id": id, "Start date": datetime(2026, 10, 2, 8),
             "resourceTypeName": "Sono Felice Country Club Mangilao",
             "Channel": "KR", "clientName": "individual", "clientGroupName": "",
             "Total": 100, "Cancelled": 0, "Creation date": datetime(2025, 1, 1)}
    value.update(extra)
    return value


class ImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("onbook", MODULE)
        cls.mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.mod)

    def aggregate(self, rows):
        return self.mod.aggregate(rows, "2026-10-07", "2026-10-01", "2026-12-31")

    def test_reference_channels_preserve_raw_and_require_same_booking(self):
        current = [row(Channel=None), row(2, Channel="JP"), row(3, Channel=None, Players=99), row(4, Channel=None)]
        reference = [row(Channel="KR"), row(2, Channel="KR"), row(3, Channel="LOCAL", Players=98)]
        enriched, audit = self.mod.reference_channels(current, reference)
        self.assertIsNone(current[0]["Channel"])
        self.assertIsNone(enriched[0]["Channel"])
        self.assertEqual(enriched[0]["reportChannel"], "KR")
        self.assertNotIn("reportChannel", enriched[1])
        self.assertNotIn("reportChannel", enriched[2])
        self.assertNotIn("reportChannel", enriched[3])
        self.assertEqual(audit["applied_rows"], 1)
        result = self.aggregate(enriched)
        self.assertEqual(result["diagnostics"]["unmapped_rows"], 2)

    def test_creation_date_does_not_exclude_preexisting_bookings(self):
        result = self.aggregate([row(), row(2, Cancelled="Yes"),
                                 row(3, **{"Start date": datetime(2027, 10, 2)}),
                                 row(4, resourceTypeName="Mangilao Restaurant Resv")])
        self.assertEqual(result["months"][0]["venues"]["mangilao"]["total"],
                         {"pax": 1, "rev": 100.0})
        self.assertEqual(result["diagnostics"]["cancelled_rows"], 1)
        self.assertEqual(result["diagnostics"]["non_golf_rows"], 1)

    def test_membership_is_not_double_counted_as_agency_or_channel(self):
        data = [row(Channel=None, clientName="J&J Marketing (MEMBER)", clientGroupName="Sono Member"),
                row(2, Channel=None, clientGroupName="Sono Member", Total=120),
                row(3, Channel=None, clientName="LAND STAR TRAVEL", Total=135)]
        v = self.aggregate(data)["months"][0]["venues"]["mangilao"]
        self.assertEqual(v["total"], {"pax": 3, "rev": 355.0})
        by_id = {x["id"]: x for x in v["categories"]}
        self.assertEqual(by_id["jj_member"]["pax"], 1)
        self.assertEqual(by_id["sono_member"]["pax"], 1)
        self.assertEqual(by_id["land_star"]["pax"], 1)

    def test_unmapped_bookings_stay_in_total(self):
        v = self.aggregate([row(Channel=None, clientName="new agency")])["months"][0]["venues"]["mangilao"]
        self.assertEqual(v["total"]["pax"], 1)
        self.assertEqual(v["categories"][-1]["id"], "unmapped")
        self.assertEqual(v["categories"][-1]["pax"], 1)

    def test_uu_nationality_uses_original_channel_and_customer_rules(self):
        data = [row(Nationality="UU", Channel="LOCAL"),
                row(2, Nationality="UU", Channel=None, clientName="HIS"),
                row(3, Nationality="UU", Channel=None, clientGroupName="Sono Member"),
                row(4, Nationality="KR", Channel="Other", clientName="DAEKUN TOUR BOOKING KOREA"),
                row(5, Nationality="UU", Channel=None, clientName="unknown agency")]
        result = self.aggregate(data)
        items = {x["id"]:x for x in result["months"][0]["venues"]["mangilao"]["categories"]}
        self.assertEqual(items["local"]["pax"],1)
        self.assertEqual(items["his"]["pax"],1)
        self.assertEqual(items["sono_member"]["pax"],1)
        self.assertEqual(items["other"]["pax"],1)
        self.assertEqual(items["unmapped"]["pax"],1)
        self.assertEqual(result["diagnostics"]["uu_golf_rows"],4)
        self.assertEqual(result["diagnostics"]["uu_unresolved_rows"],1)
        venue=result["months"][0]["venues"]["mangilao"]
        self.assertEqual(next(x for x in venue["nationalities"] if x["code"]=="UU")["pax"],4)
        self.assertEqual(next(x for x in venue["nationalities"] if x["code"]=="KR")["pax"],1)
        recovered={x["market"]:x["pax"] for x in venue["uu_resolution"]}
        self.assertEqual(recovered["LOCAL"],1)
        self.assertEqual(recovered["JP"],1)
        self.assertEqual(recovered["KR"],1)
        self.assertEqual(recovered["UNRESOLVED"],1)

    def test_original_overlapping_rules_are_preserved_and_flagged(self):
        result = self.aggregate([row(clientName="LAND STAR TRAVEL")])
        venue = result["months"][0]["venues"]["mangilao"]
        self.assertEqual(venue["total"], {"pax":2,"rev":200.0})
        self.assertEqual(venue["source_total"], {"pax":1,"rev":100.0})
        self.assertEqual(result["diagnostics"]["reconciliation_mismatches"],1)

    def test_original_total_keeps_september_october_november_scope(self):
        current = self.aggregate([row(), row(2, **{"Start date":datetime(2026,12,2)}, Total=500)])
        archive = {"month":"2026-09", "venues":{
            "mangilao":{"total":{"pax":10,"rev":1000}},
            "talofofo":{"total":{"pax":5,"rev":500}}}}
        total = self.mod.original_total(current,[archive])
        self.assertEqual(total["pax"],16)
        self.assertEqual(total["rev"],1600.0)
        self.assertEqual(total["months"],["2026-09","2026-10","2026-11"])
        self.assertIsNone(self.mod.original_total(current,[]))

    def test_duplicate_reservation_ids_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.aggregate([row(), row()])

    def test_unknown_cancel_flag_and_invalid_amount_are_rejected(self):
        for item in [row(Cancelled="unknown"), row(Total="NaN"), row(Total="bad"),
                     row(**{"Start date": None})]:
            with self.subTest(item=item), self.assertRaises(ValueError):
                self.aggregate([item])

    def test_zero_revenue_is_valid_and_blank_revenue_is_not(self):
        self.assertEqual(self.aggregate([row(Total=0)])["months"][0]["venues"]["mangilao"]["total"]["rev"], 0)
        with self.assertRaises(ValueError):
            self.aggregate([row(Total=None)])

    def test_sales_export_cannot_be_used_as_booking_export(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "sales.xlsx"
            w = openpyxl.Workbook(); w.active.append(["Id", "Production date", "Total"]); w.save(p)
            with self.assertRaisesRegex(ValueError, "reservation"):
                self.mod.load_bookings(p)

    def test_previous_snapshot_requires_matching_coverage_and_older_date(self):
        now = self.aggregate([row(), row(2, Total=150)])
        before = self.aggregate([row()]); before["as_of"] = "2026-10-06"
        self.mod.compare(now, before)
        self.assertEqual(now["months"][0]["venues"]["mangilao"]["delta"], {"pax": 1, "rev": 150.0})
        before["coverage"]["end"] = "2026-11-30"
        self.mod.compare(now, before)
        self.assertIsNone(now["comparison"])
        self.assertIsNone(now["months"][0]["venues"]["mangilao"]["delta"])

    def test_targets_preserve_blank_and_zero_and_reject_duplicate_keys(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "targets.xlsx"
            w = openpyxl.Workbook(); s = w.active
            s.append(["year", "month", "venue", "category", "budget_pax", "budget_rev", "prev_pax", "prev_rev"])
            s.append([2026, 10, "mangilao", "total", 0, None, 10, 100])
            w.save(p)
            targets = self.mod.load_targets(p)
            result = self.aggregate([row()]); self.mod.attach_targets(result, targets)
            venue = result["months"][0]["venues"]["mangilao"]
            self.assertIsNone(venue["budget_pax_rate"])
            self.assertIsNone(venue["budget_rev"])
            self.assertEqual(venue["prev_rev_change"], 0)
            s.append([2026, 10, "mangilao", "total", 5]); w.save(p)
            with self.assertRaisesRegex(ValueError, "duplicate"):
                self.mod.load_targets(p)

    def test_build_keeps_history_and_preserves_current_on_bad_import(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            p = folder / "bookings.xlsx"
            w = openpyxl.Workbook(); s = w.active
            sample = row(); s.append(list(sample)); s.append(list(sample.values())); w.save(p)
            config = {"start": "2026-10-01", "end": "2026-12-31",
                      "history": str(folder / "history"), "output": str(folder / "current.json")}
            self.mod.build(p, "2026-10-06", config)
            s.append(list(row(2, Total=150).values())); w.save(p)
            result = self.mod.build(p, "2026-10-07", config)
            self.assertEqual(result["months"][0]["venues"]["mangilao"]["delta"], {"pax": 1, "rev": 150.0})
            self.assertEqual(len(list((folder / "history").glob("*.json"))), 2)
            saved = (folder / "current.json").read_bytes()
            s.append(list(row().values())); w.save(p)
            with self.assertRaises(ValueError):
                self.mod.build(p, "2026-10-08", config)
            self.assertEqual((folder / "current.json").read_bytes(), saved)


if __name__ == "__main__":
    unittest.main()
